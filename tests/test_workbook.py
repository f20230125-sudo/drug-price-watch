"""The Excel price tracker: structure, live formulas, and the cached values Excel saved for the default inputs."""

from __future__ import annotations

import csv
import warnings

import openpyxl
import pytest

from pw.sources import MARTS, ROOT

XLSX = ROOT / "excel" / "Price-Tracker.xlsx"


@pytest.fixture(scope="module")
def wb():
    if not XLSX.exists():
        pytest.skip("workbook not built")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(XLSX, data_only=True)


def products():
    with (MARTS / "xl_products.csv").open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def test_sheets_and_names_exist(wb):
    assert wb.sheetnames == ["Read me", "Inputs", "Tracker", "Entry scenario", "Checks", "Products", "Class index", "Lists"]
    for name in ("Product", "Compare", "BaseYear", "EndYear", "ProdRow", "p_prices", "c_idx", "Keep4", "Move8"):
        assert name in wb.defined_names


def test_every_pack_label_is_unique():
    labels = [p["label"] for p in products()]
    assert len(labels) == len(set(labels)) >= 1000


def test_default_pack_prices_match_the_csv_the_verifier_reads(wb):
    p = next(x for x in products() if x["label"] == wb["Inputs"]["C4"].value)
    base, end = wb["Inputs"]["C6"].value, wb["Inputs"]["C7"].value
    assert wb["Tracker"]["B11"].value == pytest.approx(float(p[f"p{base}"]), rel=1e-12)
    assert wb["Tracker"]["B12"].value == pytest.approx(float(p[f"p{end}"]), rel=1e-12)
    assert wb["Tracker"]["B13"].value == pytest.approx(float(p[f"p{end}"]) / float(p[f"p{base}"]) - 1, rel=1e-12)
    assert wb["Checks"]["E12"].value == "ALL OK"


def test_workbook_has_live_formulas_not_pasted_results():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        f = openpyxl.load_workbook(XLSX)
    assert str(f["Tracker"]["B13"].value).startswith("=")
    assert str(f["Products"]["Z2"].value).startswith("=IF(")
    assert str(f["Entry scenario"]["D10"].value).startswith("=")


def test_the_headline_pack_is_in_the_list():
    assert any("JANUVIA 100 MG TABLET" in p["label"] for p in products())
